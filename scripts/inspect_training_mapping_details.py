from pathlib import Path

import pydicom
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parent.parent

TRAINING_DIR = PROJECT_DIR / "data" / "raw" / "training"
STUDIES_DIR = TRAINING_DIR / "Исследования"
ANNOTATIONS_FILE = TRAINING_DIR / "разметка.xlsx"


def load_excel_studies():
    workbook = load_workbook(
        ANNOTATIONS_FILE,
        read_only=True,
        data_only=True
    )

    sheet = workbook["Калибровка"]

    studies = []

    for row in sheet.iter_rows(
        min_row=3,
        values_only=True
    ):
        study = row[1]

        if study is not None:
            studies.append(str(study).strip())

    workbook.close()

    return studies


def main():
    print("=" * 80)
    print("ДЕТАЛЬНАЯ ПРОВЕРКА СВЯЗИ EXCEL ↔ ПАПКА ↔ DICOM")
    print("=" * 80)
    print()

    excel_studies = load_excel_studies()

    print(f"Исследований в Excel: {len(excel_studies)}")
    print()

    # ---------------------------------------------------------
    # Покажем первые 5 исследований из Excel
    # ---------------------------------------------------------

    print("=" * 80)
    print("ПЕРВЫЕ 5 ИССЛЕДОВАНИЙ ИЗ EXCEL")
    print("=" * 80)
    print()

    for study_uid in excel_studies[:5]:
        print(study_uid)

    print()

    # ---------------------------------------------------------
    # Находим папки исследований
    # ---------------------------------------------------------

    study_directories = [
        path
        for path in STUDIES_DIR.iterdir()
        if path.is_dir()
    ]

    print("=" * 80)
    print("ПАПКИ ИССЛЕДОВАНИЙ")
    print("=" * 80)
    print()

    print(
        f"Всего папок исследований: "
        f"{len(study_directories)}"
    )

    print()

    # ---------------------------------------------------------
    # Сравниваем имена папок с Excel
    # ---------------------------------------------------------

    excel_set = set(excel_studies)

    folder_names = {
        path.name
        for path in study_directories
    }

    folder_matches = excel_set & folder_names

    print(
        f"Совпадений Excel ↔ имя папки: "
        f"{len(folder_matches)}"
    )

    print()

    # ---------------------------------------------------------
    # Показываем подробности первых 5 папок
    # ---------------------------------------------------------

    print("=" * 80)
    print("ПОДРОБНОСТИ ПЕРВЫХ 5 ИССЛЕДОВАНИЙ")
    print("=" * 80)
    print()

    for study_dir in sorted(study_directories)[:5]:

        print("-" * 80)
        print(f"ИМЯ ПАПКИ: {study_dir.name}")

        # Все DICOM внутри этой папки
        dicom_files = sorted(
            study_dir.rglob("*.dcm")
        )

        print(
            f"Количество DICOM: "
            f"{len(dicom_files)}"
        )

        if not dicom_files:
            print("DICOM не найден.")
            print()
            continue

        # Берём первый DICOM
        file_path = dicom_files[0]

        print(
            f"Первый DICOM: "
            f"{file_path.name}"
        )

        try:
            dataset = pydicom.dcmread(
                file_path,
                stop_before_pixels=True
            )

            def get_tag(name):
                value = getattr(dataset, name, None)

                if value is None:
                    return "<ОТСУТСТВУЕТ>"

                return str(value)

            print(
                f"  StudyInstanceUID: "
                f"{get_tag('StudyInstanceUID')}"
            )

            print(
                f"  SeriesInstanceUID: "
                f"{get_tag('SeriesInstanceUID')}"
            )

            print(
                f"  SOPInstanceUID: "
                f"{get_tag('SOPInstanceUID')}"
            )

            print(
                f"  StudyID: "
                f"{get_tag('StudyID')}"
            )

            print(
                f"  SeriesNumber: "
                f"{get_tag('SeriesNumber')}"
            )

            print(
                f"  Modality: "
                f"{get_tag('Modality')}"
            )

        except Exception as error:
            print(f"Ошибка чтения DICOM: {error}")

        print()

    # ---------------------------------------------------------
    # Итог
    # ---------------------------------------------------------

    print("=" * 80)
    print("ИТОГ")
    print("=" * 80)
    print()

    if len(folder_matches) == len(excel_set):
        print(
            "Все исследования из Excel "
            "совпадают с именами папок."
        )
    else:
        print(
            "Не все исследования из Excel "
            "совпадают с именами папок."
        )

    print()

    print(
        f"Excel: {len(excel_set)}"
    )

    print(
        f"Папки: {len(folder_names)}"
    )

    print(
        f"Excel ↔ папка: {len(folder_matches)}"
    )

    print()
    print("=" * 80)


if __name__ == "__main__":
    main()