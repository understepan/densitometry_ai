from pathlib import Path
from collections import Counter

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parent.parent
TRAINING_DIR = PROJECT_DIR / "data" / "raw" / "training"
ANNOTATIONS_FILE = TRAINING_DIR / "разметка.xlsx"


def normalize_value(value):
    if value is None:
        return "<ПУСТО>"

    return str(value).strip()


def main():
    print("=" * 80)
    print("ПОЛНЫЙ АНАЛИЗ РАЗМЕТКИ")
    print("=" * 80)
    print()

    workbook = load_workbook(
        ANNOTATIONS_FILE,
        read_only=True,
        data_only=True
    )

    sheet = workbook["Калибровка"]

    # ---------------------------------------------------------
    # Названия столбцов
    # ---------------------------------------------------------

    headers_row_1 = list(
        sheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True
        )
    )[0]

    headers_row_2 = list(
        sheet.iter_rows(
            min_row=2,
            max_row=2,
            values_only=True
        )
    )[0]

    print("СТОЛБЦЫ:")
    print()

    for column_index in range(1, sheet.max_column + 1):
        header_1 = normalize_value(headers_row_1[column_index - 1])
        header_2 = normalize_value(headers_row_2[column_index - 1])

        print(
            f"{column_index:2d}. "
            f"Уровень 1: {header_1} | "
            f"Уровень 2: {header_2}"
        )

    print()

    # ---------------------------------------------------------
    # Данные
    # ---------------------------------------------------------

    data_rows = list(
        sheet.iter_rows(
            min_row=3,
            values_only=True
        )
    )

    print("=" * 80)
    print("ОБЩАЯ ИНФОРМАЦИЯ")
    print("=" * 80)

    print(f"Всего строк с данными: {len(data_rows)}")
    print(f"Всего столбцов: {sheet.max_column}")
    print()

    # ---------------------------------------------------------
    # Значения в каждом столбце
    # ---------------------------------------------------------

    print("=" * 80)
    print("УНИКАЛЬНЫЕ ЗНАЧЕНИЯ ПО СТОЛБЦАМ")
    print("=" * 80)
    print()

    for column_index in range(sheet.max_column):
        values = [
            normalize_value(row[column_index])
            for row in data_rows
        ]

        counter = Counter(values)

        header_1 = normalize_value(headers_row_1[column_index])
        header_2 = normalize_value(headers_row_2[column_index])

        print("-" * 80)
        print(
            f"Столбец {column_index + 1}: "
            f"{header_1} / {header_2}"
        )

        for value, count in counter.most_common():
            print(f"  {value}: {count}")

        print()

    # ---------------------------------------------------------
    # Комментарии
    # ---------------------------------------------------------

    print("=" * 80)
    print("КОММЕНТАРИИ")
    print("=" * 80)
    print()

    # В Excel комментарий находится в 13-м столбце
    comments_column = 12

    comments_found = []

    for row_number, row in enumerate(data_rows, start=3):
        value = row[comments_column]

        if value is not None:
            comments_found.append(
                (row_number, str(value))
            )

    print(f"Строк с комментариями: {len(comments_found)}")
    print()

    for row_number, comment in comments_found:
        print(f"Строка {row_number}: {comment}")

    print()

    workbook.close()

    print("=" * 80)
    print("АНАЛИЗ ЗАВЕРШЁН")
    print("=" * 80)


if __name__ == "__main__":
    main()