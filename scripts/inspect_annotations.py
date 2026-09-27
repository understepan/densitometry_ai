from pathlib import Path

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parent.parent
TRAINING_DIR = PROJECT_DIR / "data" / "raw" / "training"
ANNOTATIONS_FILE = TRAINING_DIR / "разметка.xlsx"


def main():
    print("=" * 70)
    print("ПРОВЕРКА ОБУЧАЮЩЕЙ РАЗМЕТКИ")
    print("=" * 70)
    print()

    if not ANNOTATIONS_FILE.exists():
        raise FileNotFoundError(
            f"Файл разметки не найден: {ANNOTATIONS_FILE}"
        )

    print(f"Файл разметки: {ANNOTATIONS_FILE}")
    print()

    workbook = load_workbook(
        ANNOTATIONS_FILE,
        read_only=True,
        data_only=True
    )

    print("Листы Excel:")

    for sheet_name in workbook.sheetnames:
        print(f"  - {sheet_name}")

    print()

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        print("=" * 70)
        print(f"ЛИСТ: {sheet_name}")
        print("=" * 70)

        print(f"Количество строк: {sheet.max_row}")
        print(f"Количество столбцов: {sheet.max_column}")
        print()

        rows = sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, 10),
            values_only=True
        )

        for row_number, row in enumerate(rows, start=1):
            print(f"Строка {row_number}:")
            print(row)
            print()

    workbook.close()

    print("=" * 70)
    print("Проверка разметки завершена.")
    print("=" * 70)


if __name__ == "__main__":
    main()