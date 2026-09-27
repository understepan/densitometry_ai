from pathlib import Path

from openpyxl import load_workbook
import pydicom


PROJECT_DIR = Path(__file__).resolve().parent.parent

TRAINING_DIR = PROJECT_DIR / "data" / "raw" / "training"
STUDIES_DIR = TRAINING_DIR / "Исследования"
ANNOTATIONS_FILE = TRAINING_DIR / "разметка.xlsx"


def load_annotation_studies():
    """
    Загружает study_id из Excel.

    В Excel study_id соответствует имени папки
    исследования, а не DICOM StudyInstanceUID.
    """

    workbook = load_workbook(
        ANNOTATIONS_FILE,
        read_only=True,
        data_only=True
    )

    sheet = workbook["Калибровка"]

    study_ids = set()

    for row in sheet.iter_rows(
        min_row=3,
        values_only=True
    ):
        study_id = row[1]

        if study_id is not None:
            study_ids.add(str(study_id).strip())

    workbook.close()

    return study_ids


def find_study_folders():
    """
    Возвращает папки исследований.
    """

    folders = [
        path
        for path in STUDIES_DIR.iterdir()
        if path.is_dir()
    ]

    return sorted(folders)


def find_dicom_files(study_folder):
    """
    Находит все DICOM внутри конкретного исследования.
    """

    return sorted(study_folder.rglob("*.dcm"))


def read_dicom_study_uid(file_path):
    """
    Читает внутренний StudyInstanceUID из DICOM.

    Он сохраняется отдельно и НЕ сравнивается
    напрямую с Excel study_id.
    """

    dataset = pydicom.dcmread(
        file_path,
        stop_before_pixels=True
    )

    return str(dataset.StudyInstanceUID)


def main():
    print("=" * 80)
    print("ПРОВЕРКА СВЯЗИ EXCEL ↔ ПАПКИ ↔ DICOM")
    print("=" * 80)
    print()

    # ---------------------------------------------------------
    # 1. Загружаем study_id из Excel
    # ---------------------------------------------------------

    annotation_studies = load_annotation_studies()

    print(
        f"Исследований в разметке Excel: "
        f"{len(annotation_studies)}"
    )
    print()

    # ---------------------------------------------------------
    # 2. Находим папки исследований
    # ---------------------------------------------------------

    study_folders = find_study_folders()

    folder_studies = {
        folder.name
        for folder in study_folders
    }

    print(
        f"Папок исследований: "
        f"{len(folder_studies)}"
    )
    print()

    # ---------------------------------------------------------
    # 3. Сравниваем Excel ↔ папки
    # ---------------------------------------------------------

    excel_and_folders = (
        annotation_studies & folder_studies
    )

    missing_folders = (
        annotation_studies - folder_studies
    )

    extra_folders = (
        folder_studies - annotation_studies
    )

    print("=" * 80)
    print("СОПОСТАВЛЕНИЕ EXCEL ↔ ПАПКИ")
    print("=" * 80)
    print()

    print(
        f"Совпало Excel ↔ папки: "
        f"{len(excel_and_folders)}"
    )

    print(
        f"Есть в Excel, но нет папки: "
        f"{len(missing_folders)}"
    )

    print(
        f"Есть папка, но нет в Excel: "
        f"{len(extra_folders)}"
    )

    print()

    if missing_folders:
        print("-" * 80)
        print("ЕСТЬ В EXCEL, НО НЕТ ПАПКИ:")
        print()

        for study_id in sorted(missing_folders):
            print(study_id)

        print()

    if extra_folders:
        print("-" * 80)
        print("ЕСТЬ ПАПКА, НО НЕТ В EXCEL:")
        print()

        for study_id in sorted(extra_folders):
            print(study_id)

        print()

    # ---------------------------------------------------------
    # 4. Проверяем DICOM внутри каждой папки
    # ---------------------------------------------------------

    print("=" * 80)
    print("ПРОВЕРКА DICOM ВНУТРИ ИССЛЕДОВАНИЙ")
    print("=" * 80)
    print()

    total_dicom = 0
    invalid_files = []
    studies_without_dicom = []

    dicom_study_uids = {}

    for study_folder in study_folders:

        study_id = study_folder.name

        dicom_files = find_dicom_files(study_folder)

        if not dicom_files:
            studies_without_dicom.append(study_id)
            continue

        total_dicom += len(dicom_files)

        study_uids_for_folder = set()

        for file_path in dicom_files:

            try:
                dicom_uid = read_dicom_study_uid(file_path)

                study_uids_for_folder.add(dicom_uid)

                if dicom_uid not in dicom_study_uids:
                    dicom_study_uids[dicom_uid] = []

                dicom_study_uids[dicom_uid].append(
                    file_path
                )

            except Exception as error:

                invalid_files.append(
                    (file_path, str(error))
                )

    print(
        f"Обработано DICOM-файлов: "
        f"{total_dicom}"
    )

    print(
        f"Уникальных DICOM StudyInstanceUID: "
        f"{len(dicom_study_uids)}"
    )

    print()

    # ---------------------------------------------------------
    # 5. Проверяем, что каждый Excel study_id имеет DICOM
    # ---------------------------------------------------------

    excel_studies_without_dicom = []

    excel_studies_with_dicom = 0

    for study_id in sorted(annotation_studies):

        study_folder = STUDIES_DIR / study_id

        if not study_folder.exists():
            excel_studies_without_dicom.append(
                study_id
            )
            continue

        dicom_files = find_dicom_files(study_folder)

        if dicom_files:
            excel_studies_with_dicom += 1
        else:
            excel_studies_without_dicom.append(
                study_id
            )

    # ---------------------------------------------------------
    # 6. Проверяем количество DICOM на исследование
    # ---------------------------------------------------------

    print("=" * 80)
    print("DICOM НА ИССЛЕДОВАНИЕ")
    print("=" * 80)
    print()

    dicom_count_distribution = {}

    for study_folder in study_folders:

        dicom_files = find_dicom_files(study_folder)

        count = len(dicom_files)

        dicom_count_distribution[count] = (
            dicom_count_distribution.get(count, 0)
            + 1
        )

    for count in sorted(dicom_count_distribution):

        studies_count = dicom_count_distribution[count]

        print(
            f"{count} DICOM: "
            f"{studies_count} исследований"
        )

    print()

    # ---------------------------------------------------------
    # 7. Показываем несколько примеров идентификаторов
    # ---------------------------------------------------------

    print("=" * 80)
    print("ПРИМЕР СВЯЗИ ID")
    print("=" * 80)
    print()

    examples_printed = 0

    for study_id in sorted(excel_and_folders):

        study_folder = STUDIES_DIR / study_id

        dicom_files = find_dicom_files(study_folder)

        if not dicom_files:
            continue

        first_dicom = dicom_files[0]

        try:
            dicom_uid = read_dicom_study_uid(
                first_dicom
            )
        except Exception:
            continue

        print(
            f"study_id (Excel/папка): {study_id}"
        )

        print(
            f"DICOM StudyInstanceUID: {dicom_uid}"
        )

        print(
            f"DICOM: {first_dicom.name}"
        )

        print()

        examples_printed += 1

        if examples_printed >= 3:
            break

    # ---------------------------------------------------------
    # 8. Ошибочные DICOM
    # ---------------------------------------------------------

    if invalid_files:

        print("=" * 80)
        print("DICOM С ОШИБКАМИ")
        print("=" * 80)
        print()

        for file_path, error in invalid_files:

            print(file_path)
            print(f"Ошибка: {error}")
            print()

    # ---------------------------------------------------------
    # 9. Исследования без DICOM
    # ---------------------------------------------------------

    if studies_without_dicom:

        print("=" * 80)
        print("ПАПКИ БЕЗ DICOM")
        print("=" * 80)
        print()

        for study_id in studies_without_dicom:
            print(study_id)

        print()

    # ---------------------------------------------------------
    # 10. Итог
    # ---------------------------------------------------------

    print("=" * 80)
    print("ИТОГ")
    print("=" * 80)
    print()

    print(
        f"Исследований в Excel: "
        f"{len(annotation_studies)}"
    )

    print(
        f"Папок исследований: "
        f"{len(folder_studies)}"
    )

    print(
        f"Совпало Excel ↔ папки: "
        f"{len(excel_and_folders)}"
    )

    print(
        f"Исследований Excel с DICOM: "
        f"{excel_studies_with_dicom}"
    )

    print(
        f"Всего DICOM: "
        f"{total_dicom}"
    )

    print(
        f"Ошибок чтения DICOM: "
        f"{len(invalid_files)}"
    )

    print()

    if (
        len(annotation_studies) == 100
        and len(folder_studies) == 100
        and len(excel_and_folders) == 100
        and excel_studies_with_dicom == 100
        and total_dicom == 499
        and not invalid_files
    ):
        print(
            "ВСЕ ПРОВЕРКИ СОПОСТАВЛЕНИЯ ПРОЙДЕНЫ."
        )
    else:
        print(
            "НАЙДЕНЫ ПРОБЛЕМЫ — НУЖНА ДОПОЛНИТЕЛЬНАЯ ПРОВЕРКА."
        )

    print()
    print("=" * 80)


if __name__ == "__main__":
    main()