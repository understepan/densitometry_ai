from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

EXCEL_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "training"
    / "разметка.xlsx"
)

OUTPUT_DIR = (
    PROJECT_ROOT
    / "results"
    / "label_analysis"
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


def normalize_value(value):
    """
    Приводит значения Excel к:
    0 / 1 / None
    """

    if value is None:
        return None

    if isinstance(value, str):

        value = value.strip()

        if value == "":
            return None

        if value == "0":
            return 0

        if value == "1":
            return 1

    if value == 0:
        return 0

    if value == 1:
        return 1

    return value


def main():

    print("=" * 70)
    print("АНАЛИЗ РАЗМЕТКИ SPINE")
    print("=" * 70)

    print()
    print("Excel:")
    print(EXCEL_PATH)

    # ========================================================
    # READ EXCEL
    # ========================================================

    wb = load_workbook(
        EXCEL_PATH,
        data_only=True
    )

    ws = wb["Калибровка"]

    rows = list(
        ws.iter_rows(
            values_only=True
        )
    )

    print()
    print(
        f"Строк Excel: {len(rows)}"
    )

    # ========================================================
    # COLUMNS
    # ========================================================
    #
    # 1  №
    # 2  study
    # 3  spine / correct positioning
    # 4  spine axis
    # 5  spine artifacts
    # ...
    # 10 final spine
    # 13 comment
    #

    records = []

    for row in rows[2:]:

        if not row:
            continue

        study = row[1]

        if study is None:
            continue

        record = {
            "study_id": str(study),

            "position": normalize_value(
                row[2]
            ),

            "axis": normalize_value(
                row[3]
            ),

            "artifacts": normalize_value(
                row[4]
            ),

            "final": normalize_value(
                row[9]
            ),

            "comment": (
                str(row[12]).strip()
                if row[12] is not None
                else ""
            ),
        }

        records.append(record)

    df = pd.DataFrame(records)

    # ========================================================
    # BASIC INFORMATION
    # ========================================================

    print()
    print("-" * 70)
    print("ОСНОВНАЯ СТАТИСТИКА")
    print("-" * 70)

    print()
    print(
        f"Исследований: {len(df)}"
    )

    for column in [
        "position",
        "axis",
        "artifacts",
        "final"
    ]:

        print()
        print(column)

        print(
            df[column]
            .value_counts(dropna=False)
            .sort_index()
        )

    # ========================================================
    # ONLY STUDIES WITH FINAL LABEL
    # ========================================================

    labeled = df[
        df["final"].isin([0, 1])
    ].copy()

    print()
    print("-" * 70)
    print("ИССЛЕДОВАНИЯ С FINAL SPINE")
    print("-" * 70)

    print(
        f"Всего: {len(labeled)}"
    )

    print(
        labeled["final"]
        .value_counts()
        .sort_index()
    )

    # ========================================================
    # FINAL VS POSITION
    # ========================================================

    print()
    print("-" * 70)
    print("FINAL vs POSITION")
    print("-" * 70)

    table = pd.crosstab(
        labeled["position"],
        labeled["final"],
        margins=True
    )

    print(table)

    # ========================================================
    # FINAL VS AXIS
    # ========================================================

    print()
    print("-" * 70)
    print("FINAL vs AXIS")
    print("-" * 70)

    table = pd.crosstab(
        labeled["axis"],
        labeled["final"],
        margins=True
    )

    print(table)

    # ========================================================
    # FINAL VS ARTIFACTS
    # ========================================================

    print()
    print("-" * 70)
    print("FINAL vs ARTIFACTS")
    print("-" * 70)

    table = pd.crosstab(
        labeled["artifacts"],
        labeled["final"],
        margins=True
    )

    print(table)

    # ========================================================
    # COMBINATION OF THREE PREVIOUS LABELS
    # ========================================================

    print()
    print("-" * 70)
    print("POSITION + AXIS + ARTIFACTS -> FINAL")
    print("-" * 70)

    combo = (
        labeled
        .groupby(
            [
                "position",
                "axis",
                "artifacts",
                "final"
            ]
        )
        .size()
        .reset_index(
            name="count"
        )
    )

    print(combo.to_string(index=False))

    # ========================================================
    # ALL THREE = 0
    # ========================================================

    clean = labeled[
        (labeled["position"] == 0)
        &
        (labeled["axis"] == 0)
        &
        (labeled["artifacts"] == 0)
    ]

    print()
    print("-" * 70)
    print("POSITION=0, AXIS=0, ARTIFACTS=0")
    print("-" * 70)

    print(
        f"Количество: {len(clean)}"
    )

    print()
    print(
        clean["final"]
        .value_counts(
            dropna=False
        )
        .sort_index()
    )

    # ========================================================
    # FINAL=1 WITH ALL THREE PREVIOUS LABELS = 0
    # ========================================================

    suspicious = clean[
        clean["final"] == 1
    ].copy()

    print()
    print("-" * 70)
    print("FINAL=1, НО POSITION=0 AXIS=0 ARTIFACTS=0")
    print("-" * 70)

    print(
        f"Количество: {len(suspicious)}"
    )

    if len(suspicious) > 0:

        print()

        print(
            suspicious[
                [
                    "study_id",
                    "position",
                    "axis",
                    "artifacts",
                    "final",
                    "comment"
                ]
            ].to_string(
                index=False
            )
        )

    # ========================================================
    # FINAL=0 WITH ANY PREVIOUS ABNORMALITY
    # ========================================================

    abnormal_previous = labeled[
        (labeled["position"] == 1)
        |
        (labeled["axis"] == 1)
        |
        (labeled["artifacts"] == 1)
    ]

    final_zero_with_abnormal = (
        abnormal_previous[
            abnormal_previous["final"] == 0
        ]
    )

    print()
    print("-" * 70)
    print("PREVIOUS LABEL=1, НО FINAL=0")
    print("-" * 70)

    print(
        f"Количество: "
        f"{len(final_zero_with_abnormal)}"
    )

    if len(final_zero_with_abnormal) > 0:

        print()

        print(
            final_zero_with_abnormal[
                [
                    "study_id",
                    "position",
                    "axis",
                    "artifacts",
                    "final",
                    "comment"
                ]
            ].to_string(
                index=False
            )
        )

    # ========================================================
    # COMMENTS
    # ========================================================

    print()
    print("-" * 70)
    print("КОММЕНТАРИИ")
    print("-" * 70)

    comments = df[
        df["comment"] != ""
    ].copy()

    print(
        f"Исследований с комментариями: "
        f"{len(comments)}"
    )

    if len(comments) > 0:

        print()

        print(
            comments[
                [
                    "study_id",
                    "final",
                    "comment"
                ]
            ].to_string(
                index=False
            )
        )

    # ========================================================
    # SAVE FULL TABLE
    # ========================================================

    csv_path = (
        OUTPUT_DIR
        / "spine_label_analysis.csv"
    )

    df.to_csv(
        csv_path,
        index=False,
        encoding="utf-8-sig"
    )

    print()
    print("-" * 70)
    print("Файл сохранён:")
    print(csv_path)
    print("-" * 70)

    # ========================================================
    # SAVE LABELED ONLY
    # ========================================================

    labeled_path = (
        OUTPUT_DIR
        / "spine_labeled_only.csv"
    )

    labeled.to_csv(
        labeled_path,
        index=False,
        encoding="utf-8-sig"
    )

    print()
    print(
        "Файл с размеченными исследованиями:"
    )

    print(labeled_path)

    print()
    print("=" * 70)
    print("АНАЛИЗ ЗАВЕРШЁН")
    print("=" * 70)


if __name__ == "__main__":
    main()