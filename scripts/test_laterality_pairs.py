from pathlib import Path
import csv

import numpy as np
import pydicom
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parent.parent

STUDIES_DIR = (
    PROJECT_DIR
    / "data"
    / "raw"
    / "training"
    / "Исследования"
)

EXCEL_FILE = (
    PROJECT_DIR
    / "data"
    / "raw"
    / "training"
    / "разметка.xlsx"
)

IMAGE_MANIFEST_FILE = (
    PROJECT_DIR
    / "data"
    / "processed"
    / "image_manifest.csv"
)


# ---------------------------------------------------------
# ПАРАМЕТРЫ
# ---------------------------------------------------------

# Пока это только кандидат.
# Мы выбрали значение между контрольными
# right (~60-71°) и left (~88-125°).
ANGLE_THRESHOLD = 80.0

# Аналогично кандидат для centroid.
CENTROID_THRESHOLD = 0.5


# ---------------------------------------------------------
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ---------------------------------------------------------

def normalize_image(image):
    image = image.astype(np.float32)

    min_value = image.min()
    max_value = image.max()

    if max_value == min_value:
        return np.zeros_like(image)

    return (
        (image - min_value)
        / (max_value - min_value)
    )


def calculate_features(image):
    """
    Вычисляет простые геометрические признаки
    изображения.
    """

    normalized = normalize_image(image)

    threshold = np.percentile(
        normalized,
        75,
    )

    mask = normalized >= threshold

    height, width = mask.shape

    ys, xs = np.where(mask)

    if len(xs) == 0:
        return None

    # -----------------------------------------------------
    # Центр яркой области
    # -----------------------------------------------------

    centroid_x = float(xs.mean())
    centroid_x_norm = centroid_x / width

    # -----------------------------------------------------
    # PCA
    # -----------------------------------------------------

    points = np.column_stack(
        (
            xs.astype(np.float64),
            ys.astype(np.float64),
        )
    )

    centered = points - points.mean(
        axis=0
    )

    covariance = np.cov(
        centered,
        rowvar=False,
    )

    eigenvalues, eigenvectors = np.linalg.eigh(
        covariance
    )

    main_vector = eigenvectors[
        :, np.argmax(eigenvalues)
    ]

    vx = float(main_vector[0])
    vy = float(main_vector[1])

    angle = np.degrees(
        np.arctan2(vy, vx)
    )

    # Приводим угол к диапазону 0...180.
    angle = angle % 180

    return {
        "centroid_x_norm": centroid_x_norm,
        "pca_angle": float(angle),
        "mask_ratio": float(mask.mean()),
    }


def get_excel_studies_with_both_hips():
    """
    Возвращает study_id, для которых Excel содержит
    разметку и правого, и левого бедра.
    """

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    sheet = workbook["Калибровка"]

    studies = set()

    # Данные начинаются после двух строк заголовка.
    for row in sheet.iter_rows(
        min_row=3,
        values_only=True,
    ):
        if not row:
            continue

        study_id = row[1]

        if study_id is None:
            continue

        # По структуре Excel:
        #
        # row[5] = right position
        # row[6] = right ROI
        # row[7] = left position
        # row[8] = left ROI

        right_values = [
            row[5],
            row[6],
        ]

        left_values = [
            row[7],
            row[8],
        ]

        has_right = any(
            value is not None
            for value in right_values
        )

        has_left = any(
            value is not None
            for value in left_values
        )

        if has_right and has_left:
            studies.add(
                str(study_id)
            )

    workbook.close()

    return studies


def get_hip_images(study_id):
    """
    Получает все hip DICOM конкретного исследования
    из image_manifest.csv.
    """

    with IMAGE_MANIFEST_FILE.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:
        rows = list(
            csv.DictReader(f)
        )

    result = []

    for row in rows:
        if row["study_id"] != study_id:
            continue

        if row["anatomy"] != "hip":
            continue

        result.append(row)

    return result


def analyze_study(study_id):
    """
    Анализирует исследование с двумя hip DICOM.
    """

    hip_images = get_hip_images(
        study_id
    )

    if len(hip_images) != 2:
        return None

    results = []

    for row in hip_images:

        dicom_path = Path(
            row["dicom_path"]
        )

        if not dicom_path.is_absolute():
            dicom_path = (
                PROJECT_DIR
                / dicom_path
            )

        if not dicom_path.exists():
            print(
                f"Файл не найден: "
                f"{dicom_path}"
            )
            continue

        ds = pydicom.dcmread(
            dicom_path
        )

        image = ds.pixel_array

        features = calculate_features(
            image
        )

        if features is None:
            continue

        centroid = (
            features["centroid_x_norm"]
        )

        angle = (
            features["pca_angle"]
        )

        if angle > ANGLE_THRESHOLD:
            angle_prediction = "left"
        else:
            angle_prediction = "right"

        if centroid > CENTROID_THRESHOLD:
            centroid_prediction = "left"
        else:
            centroid_prediction = "right"

        results.append(
            {
                "filename": row[
                    "dicom_path"
                ],
                "centroid": centroid,
                "angle": angle,
                "angle_prediction": (
                    angle_prediction
                ),
                "centroid_prediction": (
                    centroid_prediction
                ),
            }
        )

    if len(results) != 2:
        return None

    return results


# ---------------------------------------------------------
# ОСНОВНАЯ ПРОГРАММА
# ---------------------------------------------------------

def main():

    print("=" * 100)
    print(
        "ПРОВЕРКА ГЕОМЕТРИЧЕСКОГО "
        "LATERALITY НА ПАРНЫХ ИССЛЕДОВАНИЯХ"
    )
    print("=" * 100)
    print()

    studies = (
        get_excel_studies_with_both_hips()
    )

    print(
        f"Исследований с RIGHT + LEFT "
        f"в Excel: {len(studies)}"
    )

    print()

    total_pairs = 0

    angle_opposite = 0
    centroid_opposite = 0

    angle_same = 0
    centroid_same = 0

    angle_left_right_counts = 0
    centroid_left_right_counts = 0

    valid_studies = []

    for study_id in sorted(studies):

        result = analyze_study(
            study_id
        )

        if result is None:
            continue

        total_pairs += 1

        image_a = result[0]
        image_b = result[1]

        # -------------------------------------------------
        # Проверка PCA
        # -------------------------------------------------

        angle_a = image_a["angle"]
        angle_b = image_b["angle"]

        pred_a = (
            image_a["angle_prediction"]
        )

        pred_b = (
            image_b["angle_prediction"]
        )

        angle_has_left = (
            pred_a == "left"
            or pred_b == "left"
        )

        angle_has_right = (
            pred_a == "right"
            or pred_b == "right"
        )

        if (
            angle_has_left
            and angle_has_right
        ):
            angle_opposite += 1
        else:
            angle_same += 1

        # -------------------------------------------------
        # Проверка centroid
        # -------------------------------------------------

        centroid_a = (
            image_a["centroid"]
        )

        centroid_b = (
            image_b["centroid"]
        )

        centroid_pred_a = (
            image_a[
                "centroid_prediction"
            ]
        )

        centroid_pred_b = (
            image_b[
                "centroid_prediction"
            ]
        )

        centroid_has_left = (
            centroid_pred_a == "left"
            or centroid_pred_b == "left"
        )

        centroid_has_right = (
            centroid_pred_a == "right"
            or centroid_pred_b == "right"
        )

        if (
            centroid_has_left
            and centroid_has_right
        ):
            centroid_opposite += 1
        else:
            centroid_same += 1

        valid_studies.append(
            (
                study_id,
                result,
            )
        )

    # -----------------------------------------------------
    # Результаты
    # -----------------------------------------------------

    print("=" * 100)
    print("РЕЗУЛЬТАТЫ")
    print("=" * 100)
    print()

    print(
        f"Проверено пар: {total_pairs}"
    )

    print()

    print(
        "PCA angle:"
    )

    print(
        f"  пары, где алгоритм разделил "
        f"изображения на LEFT + RIGHT: "
        f"{angle_opposite}"
    )

    print(
        f"  пары, где оба изображения "
        f"получили одну сторону: "
        f"{angle_same}"
    )

    if total_pairs:
        print(
            f"  доля разделённых пар: "
            f"{angle_opposite / total_pairs * 100:.1f}%"
        )

    print()

    print(
        "Centroid X:"
    )

    print(
        f"  пары, где алгоритм разделил "
        f"изображения на LEFT + RIGHT: "
        f"{centroid_opposite}"
    )

    print(
        f"  пары, где оба изображения "
        f"получили одну сторону: "
        f"{centroid_same}"
    )

    if total_pairs:
        print(
            f"  доля разделённых пар: "
            f"{centroid_opposite / total_pairs * 100:.1f}%"
        )

    # -----------------------------------------------------
    # Подробный вывод
    # -----------------------------------------------------

    print()
    print("=" * 100)
    print(
        "ПОДРОБНО ПО КАЖДОЙ ПАРЕ"
    )
    print("=" * 100)

    for study_id, result in valid_studies:

        print()
        print(
            f"Study: {study_id}"
        )

        for i, image in enumerate(
            result,
            start=1,
        ):

            filename = Path(
                image["filename"]
            ).name

            print(
                f"  IMAGE {i}: "
                f"{filename:15s} | "
                f"centroid={image['centroid']:.4f} | "
                f"PCA={image['angle']:.2f}° | "
                f"angle_pred="
                f"{image['angle_prediction']:5s} | "
                f"centroid_pred="
                f"{image['centroid_prediction']}"
            )

    print()
    print("=" * 100)
    print("Анализ завершён.")
    print("=" * 100)


if __name__ == "__main__":
    main()